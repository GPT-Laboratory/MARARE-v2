"""Shared helpers for transcript, document, and multi-source imports."""

import base64
import os
import random
import string
import tempfile
from datetime import datetime
from typing import List, Optional, Tuple

from dotenv import load_dotenv
from langchain_core.messages import HumanMessage
from werkzeug.utils import secure_filename

from python_files.langraph_agents_hierarchy import get_openai_llm
from python_files.supabase_store import (
    get_generated_document,
    insert_generated_document,
    update_generated_document as update_generated_document_row,
)
from python_files.template_document import serialize_doc

load_dotenv()

ALLOWED_EXTENSIONS = {"txt", "pdf", "docx"}
MULTI_SOURCE_EXTENSIONS = {
    "txt", "pdf", "docx", "jpg", "jpeg", "png", "webp", "gif", "bmp",
}
IMAGE_EXTENSIONS = {"jpg", "jpeg", "png", "webp", "gif", "bmp"}
DOCUMENT_EXTENSIONS = {"txt", "pdf", "docx"}

MAX_FILE_SIZE = 15 * 1024 * 1024  # 15 MB per file
MAX_MULTI_SOURCE_TOTAL_SIZE = 30 * 1024 * 1024  # 30 MB total
MAX_MULTI_SOURCE_FILES = 15
MIN_TEXT_CHARS = 50
MAX_LLM_DOCUMENT_CHARS = 120_000
VISION_MODEL = os.getenv("VISION_MODEL", "gpt-4o-mini")


def extract_text_from_pdf(file_path: str) -> str:
    import pdfplumber

    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_path: str) -> str:
    from docx import Document

    document = Document(file_path)
    return "\n".join(
        paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
    )


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def allowed_multi_source_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in MULTI_SOURCE_EXTENSIONS


def get_file_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[1].lower() if "." in filename else ""


def classify_upload_type(filename: str) -> str:
    ext = get_file_extension(filename)
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in DOCUMENT_EXTENSIONS:
        return "document"
    return "unknown"


def extract_text_from_upload(file_storage) -> str:
    filename = secure_filename(file_storage.filename or "")
    extension = filename.rsplit(".", 1)[1].lower()

    if extension == "txt":
        raw = file_storage.read()
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                return raw.decode(encoding).strip()
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace").strip()

    with tempfile.NamedTemporaryFile(delete=False, suffix=f".{extension}") as tmp:
        file_storage.save(tmp.name)
        tmp_path = tmp.name

    try:
        if extension == "pdf":
            return extract_text_from_pdf(tmp_path).strip()
        if extension == "docx":
            return extract_text_from_docx(tmp_path).strip()
        raise ValueError(f"Unsupported file type: .{extension}")
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def extract_docx_heading_outline(file_storage) -> list:
    """Best-effort structural outline from DOCX heading styles."""
    filename = secure_filename(file_storage.filename or "")
    if not filename.lower().endswith(".docx"):
        return []

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        file_storage.seek(0)
        file_storage.save(tmp.name)
        tmp_path = tmp.name

    try:
        from docx import Document

        doc = Document(tmp_path)
        outline = []
        current_heading = None
        current_lines = []

        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue
            style_name = (para.style.name or "").lower()
            is_heading = style_name.startswith("heading") or style_name.startswith("title")

            if is_heading:
                if current_heading:
                    outline.append(
                        {
                            "heading": current_heading,
                            "content": "\n".join(current_lines).strip(),
                        }
                    )
                current_heading = text
                current_lines = []
            else:
                current_lines.append(text)

        if current_heading:
            outline.append(
                {
                    "heading": current_heading,
                    "content": "\n".join(current_lines).strip(),
                }
            )

        return outline
    except Exception as exc:
        print(f"DOCX outline extraction skipped: {exc}")
        return []
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def generate_import_meeting_id(prefix: str = "import") -> str:
    part = "".join(random.choices(string.ascii_lowercase + string.digits, k=8))
    return f"{prefix}-{part}"


def calculate_section_progress(generated_sections: dict, template_sections: list) -> float:
    if not template_sections:
        return 0.0
    filled = sum(
        1
        for section in template_sections
        if str(generated_sections.get(section.get("id"), "")).strip()
    )
    return round((filled / len(template_sections)) * 100, 1)


def save_or_update_imported_document(
    *,
    user_id: str,
    meeting_id: str,
    project_id: str,
    project_name: str,
    document_template_data: dict,
    generated_sections: dict,
    undiscussed_topics: list,
    meeting_phase: str,
    progress: float,
    use_previous_document: bool,
    existing_document_id: Optional[str],
    source: str,
):
    now = datetime.utcnow()
    version_created_at = now.isoformat()

    template_payload = {
        "categories": document_template_data.get("categories", {}),
        "sections": document_template_data.get("sections", []),
    }

    if use_previous_document and existing_document_id:
        existing = get_generated_document(existing_document_id, user_id)
        if existing:
            new_version = existing.get("version", 1) + 1
            update_data = {
                "meeting_id": meeting_id,
                "generated_sections": generated_sections,
                "undiscussed_topics": undiscussed_topics,
                "meeting_phase": meeting_phase,
                "progress": progress,
                "template": template_payload,
                "project_name": project_name,
                "source": source,
                "updated_at": now.isoformat(),
                "version": new_version,
                "version_created_at": version_created_at,
            }

            updated = update_generated_document_row(
                existing_document_id,
                user_id,
                update_data,
                append_version_history={
                    "version": new_version,
                    "createdAt": version_created_at,
                },
            )
            return serialize_doc(updated)

    initial_version = 1
    document = {
        "user_id": user_id,
        "meeting_id": meeting_id,
        "project_id": project_id,
        "project_name": project_name,
        "template": template_payload,
        "generated_sections": generated_sections,
        "undiscussed_topics": undiscussed_topics,
        "meeting_phase": meeting_phase,
        "progress": progress,
        "timestamp": version_created_at,
        "source": source,
        "created_at": now,
        "updated_at": now,
        "version": initial_version,
        "version_created_at": version_created_at,
        "version_history": [{"version": initial_version, "createdAt": version_created_at}],
    }

    saved = insert_generated_document(document)
    return serialize_doc(saved)


def _image_mime_type(extension: str) -> str:
    if extension == "jpg":
        return "jpeg"
    return extension


def extract_text_from_image_bytes(image_bytes: bytes, filename: str) -> str:
    """Extract readable text and relevant information from an image via vision model."""
    extension = get_file_extension(filename)
    mime_subtype = _image_mime_type(extension)
    b64 = base64.b64encode(image_bytes).decode("utf-8")
    data_url = f"data:image/{mime_subtype};base64,{b64}"

    prompt = (
        "Extract ALL readable text from this image. Include headings, labels, bullet points, "
        "table content, slide text, UI text, and handwritten notes if legible. "
        "Also briefly describe any diagrams, charts, or visuals that contain project-relevant facts. "
        "Return plain text only — no markdown fences."
    )

    llm = get_openai_llm(model=VISION_MODEL, temperature=0.1)
    response = llm.invoke(
        [
            HumanMessage(
                content=[
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": data_url}},
                ]
            )
        ]
    )
    text = response.content if hasattr(response, "content") else str(response)
    return (text or "").strip()


def extract_content_from_file_storage(file_storage) -> Tuple[str, str, str]:
    """
    Extract content from a single uploaded file.
    Returns (filename, source_type, extracted_text).
    """
    filename = secure_filename(file_storage.filename or "upload")
    extension = get_file_extension(filename)
    source_type = classify_upload_type(filename)

    if source_type == "image":
        raw = file_storage.read()
        content = extract_text_from_image_bytes(raw, filename)
        return filename, "image", content

    if source_type == "document":
        file_storage.seek(0)
        content = extract_text_from_upload(file_storage)
        return filename, "document", content

    raise ValueError(f"Unsupported file type: .{extension}")


def extract_all_sources(file_storages: List) -> List[dict]:
    """Extract text from multiple uploads; returns list of source dicts."""
    sources = []
    for file_storage in file_storages:
        filename, source_type, content = extract_content_from_file_storage(file_storage)
        sources.append(
            {
                "filename": filename,
                "type": source_type,
                "content": content,
                "char_count": len(content),
            }
        )
    return sources


def build_combined_source_text(sources: List[dict], max_chars: int = MAX_LLM_DOCUMENT_CHARS) -> str:
    """Merge labeled source blocks into one string for LLM mapping."""
    blocks = []
    for source in sources:
        header = f"===== SOURCE: {source['filename']} ({source['type']}) ====="
        body = (source.get("content") or "").strip()
        blocks.append(f"{header}\n{body}\n")

    combined = "\n".join(blocks).strip()
    if len(combined) > max_chars:
        combined = combined[:max_chars] + "\n\n[Content truncated for processing]"
    return combined
